import os
import re

from PyPDF2 import PdfReader
from docx import Document


# ============================================================
# KNOWN SKILLS
# ============================================================

KNOWN_SKILLS = [
    "python",
    "java",
    "c++",
    "javascript",
    "html",
    "css",
    "sql",
    "mysql",
    "sqlite",
    "flask",
    "django",
    "react",
    "node.js",
    "spring boot",
    "machine learning",
    "deep learning",
    "tensorflow",
    "pandas",
    "numpy",
    "power bi",
    "excel",
    "statistics",
    "data analysis",
    "data visualization",
    "rest api",
    "git",
    "docker",
    "figma",
    "adobe xd",
    "wireframing",
    "ui/ux",
    "communication",
    "project management",
    "data structures",
    "algorithms",
    "aws",
    "linux",
    "agile",
    "testing",
    "kubernetes"
]


# ============================================================
# KMP STRING SEARCH
# ============================================================

def build_lps(pattern):
    """
    Build the Longest Prefix Suffix (LPS) array used by KMP.
    """

    lps = [0] * len(pattern)

    length = 0
    i = 1

    while i < len(pattern):

        if pattern[i] == pattern[length]:
            length += 1
            lps[i] = length
            i += 1

        elif length != 0:
            length = lps[length - 1]

        else:
            lps[i] = 0
            i += 1

    return lps


def kmp_search(text, pattern):
    """
    Knuth-Morris-Pratt string matching.

    Returns True if pattern is found in text.
    """

    if not pattern:
        return True

    text = text.lower()
    pattern = pattern.lower()

    lps = build_lps(pattern)

    i = 0
    j = 0

    while i < len(text):

        if text[i] == pattern[j]:
            i += 1
            j += 1

            if j == len(pattern):
                return True

        else:

            if j != 0:
                j = lps[j - 1]

            else:
                i += 1

    return False


# ============================================================
# WORD BOUNDARY CHECK
# ============================================================

def contains_skill(text, skill):
    """
    Uses KMP to search for the skill while avoiding
    false matches inside larger words.

    Example:
    Java should NOT match inside JavaScript.
    """

    text_lower = text.lower()
    skill_lower = skill.lower()

    start = 0

    while start < len(text_lower):

        remaining_text = text_lower[start:]

        if not kmp_search(remaining_text, skill_lower):
            return False

        # Find the KMP match position.
        # Since KMP confirms that it exists, locate the
        # beginning using a simple scan of candidate positions.
        for position in range(
            start,
            len(text_lower) - len(skill_lower) + 1
        ):

            candidate = text_lower[
                position:position + len(skill_lower)
            ]

            if kmp_search(candidate, skill_lower):

                before_ok = (
                    position == 0
                    or not text_lower[position - 1].isalnum()
                )

                end_position = position + len(skill_lower)

                after_ok = (
                    end_position == len(text_lower)
                    or not text_lower[end_position].isalnum()
                )

                if before_ok and after_ok:
                    return True

                start = position + 1
                break

        else:
            return False

    return False


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(filepath):

    text = ""

    try:

        reader = PdfReader(filepath)

        for page in reader.pages:

            page_text = page.extract_text() or ""

            text += page_text + "\n"

    except Exception as e:

        print(
            f"[matcher] Could not parse PDF {filepath}: {e}"
        )

        return ""

    return text.strip()


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(filepath):

    text = ""

    try:

        document = Document(filepath)

        for paragraph in document.paragraphs:

            text += paragraph.text + "\n"

        # Also read text from tables
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    text += cell.text + "\n"

    except Exception as e:

        print(
            f"[matcher] Could not parse DOCX {filepath}: {e}"
        )

        return ""

    return text.strip()


# ============================================================
# GENERAL RESUME TEXT EXTRACTION
# ============================================================

def extract_text(filepath):

    extension = os.path.splitext(
        filepath
    )[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(filepath)

    elif extension == ".docx":

        return extract_text_from_docx(filepath)

    return ""


# ============================================================
# SKILL DETECTION USING KMP
# ============================================================

def detect_skills(text):

    if not text:
        return []

    found_skills = []

    for skill in KNOWN_SKILLS:

        if contains_skill(text, skill):

            found_skills.append(skill)

    return found_skills


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

def extract_education(text):

    lines = text.splitlines()

    education_keywords = [
        "education",
        "b.tech",
        "btech",
        "bachelor",
        "master",
        "degree",
        "university",
        "college",
        "cgpa",
        "gpa"
    ]

    results = []

    capture = False

    for line in lines:

        clean = line.strip()

        if not clean:
            continue

        lower = clean.lower()

        if "education" in lower:
            capture = True
            results.append(clean)
            continue

        if capture:

            if any(
                keyword in lower
                for keyword in [
                    "experience",
                    "projects",
                    "skills",
                    "certifications"
                ]
            ):
                break

            if any(
                keyword in lower
                for keyword in education_keywords
            ):
                results.append(clean)

    return " ".join(results)


# ============================================================
# EXPERIENCE EXTRACTION
# ============================================================

def extract_experience(text):

    lines = text.splitlines()

    results = []

    capture = False

    for line in lines:

        clean = line.strip()

        if not clean:
            continue

        lower = clean.lower()

        if "experience" in lower:

            capture = True
            results.append(clean)
            continue

        if capture:

            if any(
                keyword in lower
                for keyword in [
                    "education",
                    "projects",
                    "skills",
                    "certifications"
                ]
            ):
                break

            results.append(clean)

    return " ".join(results)


# ============================================================
# CERTIFICATION EXTRACTION
# ============================================================

def extract_certifications(text):

    lines = text.splitlines()

    results = []

    capture = False

    for line in lines:

        clean = line.strip()

        if not clean:
            continue

        lower = clean.lower()

        if "certification" in lower:

            capture = True
            results.append(clean)
            continue

        if capture:

            if any(
                keyword in lower
                for keyword in [
                    "education",
                    "experience",
                    "projects",
                    "skills"
                ]
            ):
                break

            results.append(clean)

    return " ".join(results)


# ============================================================
# COMPLETE RESUME DETAILS
# ============================================================

def extract_resume_details(text):

    skills = detect_skills(text)

    education = extract_education(text)

    experience = extract_experience(text)

    certifications = extract_certifications(text)

    return {
        "skills": skills,
        "education": education,
        "experience": experience,
        "certifications": certifications
    }


# ============================================================
# JOB MATCHING
# ============================================================

def compute_match_score(
    resume_skills,
    job_skills_csv
):

    job_skills = [
        skill.strip().lower()
        for skill in job_skills_csv.split(",")
        if skill.strip()
    ]

    resume_skills_lower = [
        skill.strip().lower()
        for skill in resume_skills
    ]

    if not job_skills:

        return 0.0, [], []

    matched = []

    missing = []

    for skill in job_skills:

        if skill in resume_skills_lower:

            matched.append(skill)

        else:

            missing.append(skill)

    score = round(
        (len(matched) / len(job_skills)) * 100,
        1
    )

    return score, matched, missing


# ============================================================
# TEST
# ============================================================

if __name__ == "__main__":

    print("========================================")
    print("SKILLBRIDGE MATCHER")
    print("========================================")
    print("Algorithm: KMP String Matching")
    print("Matching: Rule-Based Skill Comparison")
    print("========================================")